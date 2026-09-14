import shutil
import subprocess

import pytest
from docx import Document
from docx.oxml.ns import qn

from docx_normalizer import normalize, normalize_file
from docx_normalizer.config import load_config
from docx_normalizer.xmlutil import P, block_list, has_ppr_flag, local, para_text
from make_fixtures import build_broken

CFG = load_config(overrides={"toc": {"refresh_with_libreoffice": False}})

EXPECTED_HEADINGS = {
    "1 Introduction": "Heading 1",
    "1.1 Purpose": "Heading 2",
    "2 Scope": "Heading 1",
    "2.1 In scope": "Heading 2",
    "3 Fees": "Heading 1",
    "Appendix A Glossary": "Heading 1",
}
NOT_HEADINGS = [
    "ACME Supplier Policy",
    "This is a bold sentence that should stay body text.",
    "2. Short item",
    "Fees are listed below:",
    "Table 2: Contacts",
]


def run(tmp_path, kind):
    src = build_broken(tmp_path / f"broken_{kind}.docx", kind)
    doc = Document(str(src))
    report = normalize(doc, CFG)
    out = tmp_path / f"out_{kind}.docx"
    doc.save(str(out))
    return Document(str(out)), report


def by_text(doc):
    return {p.text.strip(): p for p in doc.paragraphs if p.text.strip()}


@pytest.fixture(params=["hand", "mixed"])
def result(request, tmp_path):
    return run(tmp_path, request.param)


def test_headings_detected_with_correct_levels(result):
    doc, _ = result
    paras = by_text(doc)
    for text, style in EXPECTED_HEADINGS.items():
        assert paras[text].style.name == style, text


def test_traps_are_not_promoted(result):
    doc, report = result
    paras = by_text(doc)
    for text in NOT_HEADINGS:
        assert not paras[text].style.name.startswith("Heading"), text
    assert doc.tables[0].cell(1, 0).paragraphs[0].style.name == "Normal"
    assert "2. Short item" in [h.text for h in report.near_misses]


def test_text_box_unwrapped_once(result):
    doc, _ = result
    body = doc.element.body
    assert next(body.iter(qn("w:txbxContent")), None) is None
    assert sum(1 for p in body.iter(P) if para_text(p).strip() == "3 Fees") == 1


def test_old_toc_replaced_by_single_field(result):
    doc, report = result
    body = doc.element.body
    instr = [i.text for i in body.iter(qn("w:instrText")) if "TOC" in (i.text or "")]
    assert len(instr) == 1
    texts = [para_text(p).strip() for p in body.iter(P)]
    assert not any("...." in t for t in texts)
    assert texts.count("Table of Contents") == 1
    assert "Contents" not in texts
    assert "Purpose" not in " ".join(t for t in texts if "\t" in t)
    assert report.toc_inserted
    assert "hand-typed" in report.toc_found
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_page_breaks(result):
    doc, _ = result
    paras = by_text(doc)
    for text, style in EXPECTED_HEADINGS.items():
        assert bool(paras[text].paragraph_format.page_break_before) == (style == "Heading 1"), text
    assert paras["Table of Contents"].paragraph_format.page_break_before
    # Table 1: its lead-in chain starts with "3 Fees", which already breaks.
    assert not paras["Fees are listed below:"].paragraph_format.page_break_before
    assert paras["Fees are listed below:"].paragraph_format.keep_with_next
    # Table 2: caption moves to the new page with it.
    assert paras["Table 2: Contacts"].paragraph_format.page_break_before
    # Table 3: no lead-in -> dedicated break paragraph right before it.
    t3 = doc.tables[2]._tbl
    assert has_ppr_flag(t3.getprevious(), "pageBreakBefore")
    # Manual breaks are gone.
    assert not [br for br in doc.element.body.iter(qn("w:br")) if br.get(qn("w:type")) == "page"]


def _signature(doc):
    sig = []
    for b in block_list(doc.element.body):
        if local(b) == "tbl":
            sig.append(("tbl", len(b.findall(qn("w:tr")))))
        else:
            ppr = b.find(qn("w:pPr"))
            ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
            sig.append((ps.get(qn("w:val")) if ps is not None else None, para_text(b), has_ppr_flag(b, "pageBreakBefore")))
    return sig


@pytest.mark.parametrize("kind", ["hand", "mixed"])
def test_idempotent(tmp_path, kind):
    doc, _ = run(tmp_path, kind)
    first = _signature(doc)
    normalize(doc, CFG)
    assert _signature(doc) == first


def test_unknown_config_key_rejected():
    with pytest.raises(ValueError):
        load_config(overrides={"headings": {"max_wrds": 3}})


def _libreoffice_available():
    if shutil.which("soffice") is None:
        return False
    cfg = load_config()
    try:
        return subprocess.run([cfg["toc"]["uno_python"], "-c", "import uno"], capture_output=True).returncode == 0
    except FileNotFoundError:
        return False


@pytest.mark.skipif(not _libreoffice_available(), reason="LibreOffice + python uno not available")
def test_libreoffice_fills_toc(tmp_path):
    src = build_broken(tmp_path / "broken.docx", "hand")
    out = tmp_path / "out.docx"
    report = normalize_file(src, out, load_config())
    assert report.toc_refreshed.startswith("yes"), report.warnings
    texts = [para_text(p) for p in Document(str(out)).element.body.iter(P)]
    assert sum("Purpose" in t for t in texts) >= 2  # TOC entry + heading
    assert not any("Right-click" in t for t in texts)
